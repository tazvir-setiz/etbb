from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Person, Survey, SurveyPerson, User, Vote


async def build_survey_workbook(session: AsyncSession, survey_id: int) -> BytesIO:
    survey = await session.get(Survey, survey_id)
    if survey is None:
        raise ValueError("نظرسنجی پیدا نشد.")
    votes = await session.execute(
        select(Vote, Person.display_name, User.full_name, User.platform_user_id, User.platform)
        .join(Person, Person.id == Vote.person_id)
        .join(User, User.id == Vote.user_id)
        .where(Vote.survey_id == survey_id)
        .order_by(Person.display_name, Vote.created_at)
    )
    participant_stats = await session.execute(
        select(Person.display_name, func.count(Vote.id), func.avg(Vote.score))
        .join(SurveyPerson, SurveyPerson.person_id == Person.id)
        .outerjoin(Vote, (Vote.person_id == Person.id) & (Vote.survey_id == survey_id))
        .where(SurveyPerson.survey_id == survey_id)
        .group_by(Person.id, Person.display_name)
        .order_by(Person.display_name)
    )
    voter_count = await session.scalar(
        select(func.count(func.distinct(Vote.user_id))).where(Vote.survey_id == survey_id)
    ) or 0
    workbook = Workbook()
    summary = workbook.active
    summary.title = "خلاصه"
    summary.append(["فرد", "تعداد امتیاز", "میانگین"])
    detail = workbook.create_sheet("امتیازها")
    detail.append(["فرد", "امتیاز", "نظر", "زمان ثبت"] + ([] if survey.anonymous else ["رأی‌دهنده", "شناسه کاربر", "پیام‌رسان"]))
    progress = workbook.create_sheet("پیشرفت")
    progress.append(["شاخص", "مقدار"])
    participant_rows = list(participant_stats)
    rated_people = sum(1 for _, count, _ in participant_rows if count)
    total_votes = sum(count for _, count, _ in participant_rows)
    progress.append(["افراد تخصیص‌یافته", len(participant_rows)])
    progress.append(["افراد دارای رأی", rated_people])
    progress.append(["پیشرفت افراد (درصد)", round(rated_people / len(participant_rows) * 100, 1) if participant_rows else 0])
    progress.append(["کل رأی‌ها", total_votes])
    progress.append(["رأی‌دهندگان یکتا", voter_count])
    for name, count, average in participant_rows:
        summary.append([name, count, float(average or 0)])
    for vote, person_name, voter_name, voter_id, platform in votes:
        record: list[object] = [person_name, vote.score, vote.comment or "", vote.created_at]
        if not survey.anonymous:
            record.extend([voter_name, voter_id, platform])
        detail.append(record)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="right")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 32)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
