import csv
import logging
from io import BytesIO, StringIO
from itertools import islice
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

logger = logging.getLogger(__name__)

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000

_HEADERS = {
    "internal_id": {"internal_id", "id", "شناسه داخلی", "شناسهٔ داخلی"},
    "first_name": {"first_name", "نام"},
    "last_name": {"last_name", "نام خانوادگی", "نام‌خانوادگی"},
    "description": {"description", "توضیح", "توضیحات"},
}


class PeopleImportError(ValueError):
    """The uploaded file cannot be read as a people import."""


def _column_map(headers: list[object]) -> dict[str, int]:
    normalized = [str(value or "").strip().casefold() for value in headers]
    mapping: dict[str, int] = {}
    for field, aliases in _HEADERS.items():
        for index, value in enumerate(normalized):
            if value in {alias.casefold() for alias in aliases}:
                mapping[field] = index
                break
    required = {"internal_id", "first_name", "last_name"}
    if not required <= mapping.keys():
        raise PeopleImportError("ستون‌های الزامی: شناسه داخلی، نام و نام خانوادگی")
    return mapping


def _records_from_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    if not rows:
        raise PeopleImportError("فایل خالی است.")
    mapping = _column_map(rows[0])
    result: list[dict[str, object]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if row_number > MAX_IMPORT_ROWS + 1:
            raise PeopleImportError(f"حداکثر {MAX_IMPORT_ROWS} ردیف داده پذیرفته می‌شود.")
        if not any(str(value or "").strip() for value in values):
            continue
        def cell(field: str) -> str:
            index = mapping.get(field)
            return str(values[index] or "").strip() if index is not None and index < len(values) else ""

        record = {
            "row": row_number,
            "internal_id": cell("internal_id"),
            "first_name": cell("first_name"),
            "last_name": cell("last_name"),
            "description": cell("description") or None,
        }
        errors = []
        if not record["internal_id"] or len(record["internal_id"]) > 64:
            errors.append("شناسهٔ داخلی باید ۱ تا ۶۴ نویسه باشد")
        if not record["first_name"] or len(record["first_name"]) > 100:
            errors.append("نام باید ۱ تا ۱۰۰ نویسه باشد")
        if not record["last_name"] or len(record["last_name"]) > 100:
            errors.append("نام خانوادگی باید ۱ تا ۱۰۰ نویسه باشد")
        if record["description"] and len(record["description"]) > 2000:
            errors.append("توضیح حداکثر ۲۰۰۰ نویسه است")
        record["error"] = "؛ ".join(errors) or None
        result.append(record)
    if not result:
        raise PeopleImportError("فایل ردیف داده ندارد.")
    return result


def parse_people_file(filename: str, content: bytes) -> list[dict[str, object]]:
    logger.info("Parsing people import file '%s' (size=%d bytes)", filename, len(content))
    if len(content) > MAX_IMPORT_BYTES:
        logger.warning("People import rejected: file '%s' exceeds 5MB size limit", filename)
        raise PeopleImportError("حجم فایل نباید بیشتر از ۵ مگابایت باشد.")
    suffix = Path(filename).suffix.casefold()
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            sample = text[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(StringIO(text), dialect))
        elif suffix in {".xlsx", ".xlsm"}:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            rows = list(islice(workbook.active.iter_rows(values_only=True), MAX_IMPORT_ROWS + 2))
            workbook.close()
            if len(rows) > MAX_IMPORT_ROWS + 1:
                logger.warning("People import rejected: '%s' has more than %d rows", filename, MAX_IMPORT_ROWS)
                raise PeopleImportError(f"حداکثر {MAX_IMPORT_ROWS} ردیف داده پذیرفته می‌شود.")
        else:
            logger.warning("People import rejected: unsupported file suffix '%s'", suffix)
            raise PeopleImportError("فقط فایل CSV یا XLSX پذیرفته می‌شود.")
    except (UnicodeDecodeError, csv.Error, OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        if isinstance(exc, PeopleImportError):
            raise
        logger.warning("People import read error for '%s': %s", filename, exc)
        raise PeopleImportError("خواندن فایل ممکن نشد؛ قالب و encoding آن را بررسی کنید.") from exc

    records = _records_from_rows([list(row) for row in rows])
    valid_count = sum(1 for r in records if not r.get("error"))
    invalid_count = len(records) - valid_count
    logger.info(
        "Successfully parsed people import file '%s': total=%d, valid=%d, invalid=%d",
        filename,
        len(records),
        valid_count,
        invalid_count,
    )
    return records
