from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.exports.excel import build_survey_workbook
from app.models import Base, Person, Survey, SurveyPerson, SurveyStatus
from app.services.voting import VotingError, record_vote, survey_statistics, survey_status


@pytest.fixture
async def session_factory(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'test_quality.db'}")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    yield factory
    await engine.dispose()


@pytest.mark.asyncio
async def test_time_window_exact_boundaries(session_factory):
    now = datetime(2026, 10, 11, 12, 0, 0, tzinfo=timezone.utc)
    start = now
    end = now + timedelta(hours=2)

    async with session_factory() as session:
        person = Person(first_name="مهدی", last_name="کاظمی", display_name="مهدی کاظمی", internal_id="M-1")
        survey = Survey(
            title="ارزیابی مرزی",
            starts_at=start,
            ends_at=end,
            status=SurveyStatus.ACTIVE.value,
            allow_vote_editing=True,
        )
        session.add_all([person, survey])
        await session.flush()
        session.add(SurveyPerson(survey_id=survey.id, person_id=person.id))
        await session.commit()
        s_id, p_id = survey.id, person.id

    # Exactly 1 second before start -> rejected
    async with session_factory() as session:
        with pytest.raises(VotingError, match="فعال نیست"):
            await record_vote(
                session,
                platform="telegram",
                platform_user_id=10,
                survey_id=s_id,
                person_id=p_id,
                score=4,
                now=start - timedelta(seconds=1),
            )

    # Exactly at start -> accepted
    async with session_factory() as session:
        vote_start = await record_vote(
            session,
            platform="telegram",
            platform_user_id=10,
            survey_id=s_id,
            person_id=p_id,
            score=4,
            now=start,
        )
        assert vote_start.score == 4

    # Exactly at end -> accepted
    async with session_factory() as session:
        vote_end = await record_vote(
            session,
            platform="telegram",
            platform_user_id=11,
            survey_id=s_id,
            person_id=p_id,
            score=5,
            now=end,
        )
        assert vote_end.score == 5

    # 1 second after end -> rejected
    async with session_factory() as session:
        with pytest.raises(VotingError, match="فعال نیست"):
            await record_vote(
                session,
                platform="telegram",
                platform_user_id=12,
                survey_id=s_id,
                person_id=p_id,
                score=3,
                now=end + timedelta(seconds=1),
            )


@pytest.mark.asyncio
async def test_vote_editing_permission_and_comment_policy(session_factory):
    now = datetime(2026, 10, 11, 12, 0, 0, tzinfo=timezone.utc)
    async with session_factory() as session:
        person = Person(first_name="سارا", last_name="کریمی", display_name="سارا کریمی", internal_id="S-1")
        # Survey with editing disabled and comments disabled
        survey = Survey(
            title="نظرسنجی بدون ویرایش و نظر",
            starts_at=now - timedelta(hours=1),
            ends_at=now + timedelta(hours=1),
            status=SurveyStatus.ACTIVE.value,
            allow_vote_editing=False,
            allow_comments=False,
            min_score=1,
            max_score=5,
        )
        session.add_all([person, survey])
        await session.flush()
        session.add(SurveyPerson(survey_id=survey.id, person_id=person.id))
        await session.commit()
        s_id, p_id = survey.id, person.id

    # First vote (comment must be ignored because allow_comments=False)
    async with session_factory() as session:
        vote1 = await record_vote(
            session,
            platform="telegram",
            platform_user_id=20,
            survey_id=s_id,
            person_id=p_id,
            score=4,
            comment="نظر نامعتبر",
            now=now,
        )
        assert vote1.score == 4
        assert vote1.comment is None

    # Second vote must be rejected because allow_vote_editing=False
    async with session_factory() as session:
        with pytest.raises(VotingError, match="قبلاً به این فرد امتیاز داده‌اید"):
            await record_vote(
                session,
                platform="telegram",
                platform_user_id=20,
                survey_id=s_id,
                person_id=p_id,
                score=5,
                now=now,
            )


@pytest.mark.asyncio
async def test_cancelled_survey_preserves_stats_and_rejects_votes(session_factory):
    now = datetime(2026, 10, 11, 12, 0, 0, tzinfo=timezone.utc)
    async with session_factory() as session:
        person = Person(first_name="حسین", last_name="حسینی", display_name="حسین حسینی", internal_id="H-1")
        survey = Survey(
            title="نظرسنجی لغوشونده",
            starts_at=now - timedelta(hours=1),
            ends_at=now + timedelta(hours=1),
            status=SurveyStatus.ACTIVE.value,
            allow_vote_editing=True,
        )
        session.add_all([person, survey])
        await session.flush()
        session.add(SurveyPerson(survey_id=survey.id, person_id=person.id))
        await session.commit()
        s_id, p_id = survey.id, person.id

    # Record a vote while active
    async with session_factory() as session:
        await record_vote(
            session,
            platform="telegram",
            platform_user_id=30,
            survey_id=s_id,
            person_id=p_id,
            score=4,
            now=now,
        )

    # Cancel survey
    async with session_factory() as session:
        survey = await session.get(Survey, s_id)
        survey.status = SurveyStatus.CANCELLED.value
        await session.commit()

    # Voting now must be rejected
    async with session_factory() as session:
        survey = await session.get(Survey, s_id)
        assert survey_status(survey, now) == SurveyStatus.CANCELLED
        with pytest.raises(VotingError, match="فعال نیست"):
            await record_vote(
                session,
                platform="telegram",
                platform_user_id=31,
                survey_id=s_id,
                person_id=p_id,
                score=5,
                now=now,
            )

    # Statistics and previous votes must still be accurate
    async with session_factory() as session:
        stats = await survey_statistics(session, s_id)
        assert len(stats) == 1
        assert stats[0]["votes"] == 1
        assert stats[0]["average"] == 4.0


@pytest.mark.asyncio
async def test_survey_anonymity_in_excel_export(session_factory):
    now = datetime(2026, 10, 11, 12, 0, 0, tzinfo=timezone.utc)

    # 1. Anonymous survey
    async with session_factory() as session:
        person = Person(first_name="پیمان", last_name="کریمی", display_name="پیمان کریمی", internal_id="P-99")
        survey = Survey(
            title="نظرسنجی ناشناس",
            starts_at=now - timedelta(hours=1),
            ends_at=now + timedelta(hours=1),
            status=SurveyStatus.ACTIVE.value,
            anonymous=True,
        )
        session.add_all([person, survey])
        await session.flush()
        session.add(SurveyPerson(survey_id=survey.id, person_id=person.id))
        await session.commit()
        s_anon_id, p_id = survey.id, person.id

        await record_vote(
            session,
            platform="telegram",
            platform_user_id=40,
            survey_id=s_anon_id,
            person_id=p_id,
            score=5,
            now=now,
        )

    async with session_factory() as session:
        stream = await build_survey_workbook(session, s_anon_id)
        wb = load_workbook(stream)
        assert "امتیازها" in wb.sheetnames
        headers = [cell.value for cell in wb["امتیازها"][1]]
        # Must not contain individual voter details columns
        assert "رأی‌دهنده" not in headers
        assert "شناسه کاربر" not in headers
        assert "پیام‌رسان" not in headers
        wb.close()

    # 2. Non-anonymous survey
    async with session_factory() as session:
        person2 = Person(first_name="آرش", last_name="فرهادی", display_name="آرش فرهادی", internal_id="A-99")
        survey2 = Survey(
            title="نظرسنجی آشکار",
            starts_at=now - timedelta(hours=1),
            ends_at=now + timedelta(hours=1),
            status=SurveyStatus.ACTIVE.value,
            anonymous=False,
        )
        session.add_all([person2, survey2])
        await session.flush()
        session.add(SurveyPerson(survey_id=survey2.id, person_id=person2.id))
        await session.commit()
        s_public_id, p2_id = survey2.id, person2.id

        await record_vote(
            session,
            platform="telegram",
            platform_user_id=50,
            survey_id=s_public_id,
            person_id=p2_id,
            score=4,
            now=now,
        )

    async with session_factory() as session:
        stream2 = await build_survey_workbook(session, s_public_id)
        wb2 = load_workbook(stream2)
        assert "امتیازها" in wb2.sheetnames
        headers2 = [cell.value for cell in wb2["امتیازها"][1]]
        # Must contain voter details columns for non-anonymous survey
        assert "رأی‌دهنده" in headers2
        assert "شناسه کاربر" in headers2
        assert "پیام‌رسان" in headers2
        wb2.close()


def test_dual_platform_bot_configuration():
    from aiogram import Bot
    from aiogram.client.session.aiohttp import AiohttpSession
    from aiogram.client.telegram import TelegramAPIServer

    dummy_token = "123456:ABC-DEF1234ghIkl-zyx57W2v1u123ew11"

    # Telegram bot
    bot_tg = Bot(token=dummy_token)
    assert "telegram.org" in bot_tg.session.api.base

    # Bale bot
    bale_session = AiohttpSession(api=TelegramAPIServer.from_base("https://tapi.bale.ai"))
    bot_bale = Bot(token=dummy_token, session=bale_session)
    assert "tapi.bale.ai" in bot_bale.session.api.base

